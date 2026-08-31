from __future__ import annotations

from openpyxl.chart import BarChart
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.data_source import AxDataSource, NumData, NumDataSource, NumVal, StrData, StrVal
from openpyxl.chart.legend import Legend
from openpyxl.chart.series import Series, SeriesLabel
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

CHART_COL_WIDTH = 18
CHART_HEIGHT = 12

GroupedSeries = tuple[str, list[float]]


def _configure_value_axis(chart: BarChart) -> None:
    chart.y_axis.axPos = "l"
    chart.y_axis.tickLblPos = "nextTo"
    chart.y_axis.majorTickMark = "in"
    chart.y_axis.delete = False
    chart.y_axis.majorGridlines = ChartLines()
    chart.y_axis.title = None

    chart.x_axis.axPos = "b"
    chart.x_axis.tickLblPos = "low"
    chart.x_axis.majorTickMark = "in"
    chart.x_axis.delete = False
    chart.x_axis.title = None


def _str_data(labels: list[str]) -> StrData:
    pts = [StrVal(idx=i, v=str(label)) for i, label in enumerate(labels)]
    return StrData(ptCount=len(pts), pt=pts)


def _num_data(values: list[float]) -> NumData:
    pts = [NumVal(idx=i, v=v) for i, v in enumerate(values)]
    return NumData(ptCount=len(pts), pt=pts)


def add_grouped_bar_chart(
    ws: Worksheet,
    anchor: str,
    *,
    title: str,
    category_labels: list[str],
    series: list[GroupedSeries],
) -> BarChart:
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "clustered"
    chart.overlap = 0
    chart.gapWidth = 80
    chart.style = 10
    chart.title = title
    chart.height = CHART_HEIGHT
    chart.width = CHART_COL_WIDTH
    chart.legend = Legend()
    chart.legend.position = "r"
    _configure_value_axis(chart)

    categories = _str_data(category_labels)
    for idx, (seq_label, y_values) in enumerate(series):
        s = Series()
        s.title = SeriesLabel(v=seq_label)
        s.val = NumDataSource(numLit=_num_data(y_values))
        if idx == 0:
            s.cat = AxDataSource(strLit=categories)
        chart.series.append(s)

    ws.add_chart(chart, anchor)
    return chart


def write_chart_row(
    ws: Worksheet,
    anchor_row: int,
    charts: list[tuple[str, list[str], list[GroupedSeries]]],
) -> int:
    for idx, (title, category_labels, series) in enumerate(charts):
        if not category_labels or not series:
            continue
        anchor = f"{get_column_letter(1 + idx * CHART_COL_WIDTH)}{anchor_row}"
        add_grouped_bar_chart(
            ws,
            anchor,
            title=title,
            category_labels=category_labels,
            series=series,
        )
    return anchor_row + 1
