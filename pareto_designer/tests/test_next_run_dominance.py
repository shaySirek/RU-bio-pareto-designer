from __future__ import annotations

import math
from pathlib import Path

import numpy as np

from pareto_designer.models.context import ParetoResult
from pareto_designer.shared.seq_design_utils.solution_quality import RoiDistribution
from pareto_designer.views.experiment_report.dominance import (
    attach_next_run_dominance,
    n_dominated_by,
    sweep_dominance,
)
from pareto_designer.views.experiment_report.excel_schema import design_run_table
from pareto_designer.views.experiment_report.metrics import (
    build_design_run_summaries,
    design_run_summary,
    sort_design_runs,
)
from pareto_designer.views.experiment_report.models import (
    LoadedRun,
    RunParams,
    SamplerParams,
)
from pareto_designer.views.experiment_report.sweeps import set_sweep_membership


def _sol(sid, cost, binding, *, hits=None):
    return ParetoResult(
        cost=cost,
        binding_score=binding,
        origin_binding_score=binding,
        id=sid,
        url=f"{sid}_details.html",
        txt_file=f"{sid}.txt",
        fasta_file=f"{sid}.fa",
        positional_objectives_file=f"{sid}.npy",
        max_positional_cost=1.0,
        min_positional_binding=0.0,
        max_positional_binding=1.0,
        sequence="ACGT",
        n_cost_items=1,
        motif_hits=hits or [],
    )


def _run(
    seq_id,
    *,
    solutions,
    sweeps,
    k=100,
    alpha=1.0,
    log_pos=False,
    reduce=0.875,
    fsm=2048,
):
    fsm_id = "logexp_db_fsm" if reduce == 0 else f"logexp_reduced_fsm_{fsm}"
    run = LoadedRun(
        params=RunParams(
            seq_id=seq_id,
            fsm_id=fsm_id,
            fsm_size=fsm,
            reduce_fsm_by=reduce,
            sampler=SamplerParams(k=k, alpha=alpha, log_pos=log_pos),
        ),
        metadata={
            "n_solutions": len(solutions),
            "runtime": "0:00:01",
            "fsm_binding_score_err": 0.1,
            "db_fsm_size": 8192,
        },
        solutions=solutions,
        path=Path("dummy/results_metadata.json"),
    )
    set_sweep_membership(run, sweeps)
    return run


def _summary(run, sweeps):
    return design_run_summary(run, sweeps, w=500.0)


def test_min_min_dominance():
    assert n_dominated_by([_sol("a", 10, 5)], [_sol("b", 10, 5)]) == 0
    assert n_dominated_by([_sol("a", 10, 8), _sol("b", 20, 3)], [_sol("c", 9, 8)]) == 1


def _assert_empty_dist(dist: RoiDistribution) -> None:
    assert all(
        math.isnan(getattr(dist, stat))
        for stat in ("min", "p25", "p50", "p75", "max", "mean", "std")
    )


def _assert_dist(dist: RoiDistribution, values: list[float]) -> None:
    arr = np.asarray(values, dtype=float)
    assert dist.min == float(np.min(arr))
    assert dist.p25 == float(np.percentile(arr, 25))
    assert dist.p50 == float(np.percentile(arr, 50))
    assert dist.p75 == float(np.percentile(arr, 75))
    assert dist.max == float(np.max(arr))
    assert dist.mean == float(np.mean(arr))
    assert dist.std == float(np.std(arr))


def test_region_counts_and_missing_next():
    a = _run(
        "s",
        solutions=[_sol("h", 10, 20, hits=[(1, 7)]), _sol("r", 20, 5)],
        sweeps=["alpha"],
    )
    b = _run("s", alpha=2.0, solutions=[_sol("n", 9, 19)], sweeps=["alpha"])
    stats = sweep_dominance(a, b, w=500.0)
    assert stats.n_by_region == {"hits": 1, "nonsyn": 0, "roi": 0, "plateau": 0}
    assert stats.n_global == 1
    _assert_dist(stats.func_cost_gain_by_region["hits"], [1.0])
    _assert_dist(stats.binding_gain_by_region["hits"], [1.0])
    _assert_dist(stats.func_cost_gain, [1.0])
    _assert_dist(stats.binding_gain, [1.0])
    _assert_empty_dist(stats.func_cost_gain_by_region["roi"])
    _assert_empty_dist(stats.binding_gain_by_region["nonsyn"])

    missing = sweep_dominance(a, None)
    assert missing.n_by_region["hits"] is None
    assert missing.n_global is None
    _assert_empty_dist(missing.func_cost_gain)
    _assert_empty_dist(missing.binding_gain)
    _assert_empty_dist(missing.func_cost_gain_by_region["hits"])


def test_tightest_partner_gains():
    a = _run(
        "s",
        solutions=[_sol("p", 10, 8, hits=[(1, 7)])],
        sweeps=["alpha"],
    )
    b = _run(
        "s",
        alpha=2.0,
        solutions=[_sol("near", 9, 8), _sol("far", 1, 1)],
        sweeps=["alpha"],
    )
    stats = sweep_dominance(a, b, w=500.0)
    assert stats.n_by_region["hits"] == 1
    assert stats.n_global == 1
    _assert_dist(stats.func_cost_gain, [1.0])
    _assert_dist(stats.binding_gain, [0.0])
    _assert_dist(stats.func_cost_gain_by_region["hits"], [1.0])
    _assert_dist(stats.binding_gain_by_region["hits"], [0.0])


def test_global_gains_cover_all_regions():
    a = _run(
        "s",
        solutions=[_sol("h", 10, 20, hits=[(1, 7)]), _sol("r", 20, 5)],
        sweeps=["alpha"],
    )
    b = _run(
        "s",
        alpha=2.0,
        solutions=[_sol("qh", 9, 19), _sol("qr", 18, 4)],
        sweeps=["alpha"],
    )
    stats = sweep_dominance(a, b, w=500.0)
    assert stats.n_by_region == {"hits": 1, "nonsyn": 0, "roi": 1, "plateau": 0}
    assert stats.n_global == 2
    _assert_dist(stats.func_cost_gain_by_region["hits"], [1.0])
    _assert_dist(stats.binding_gain_by_region["hits"], [1.0])
    _assert_dist(stats.func_cost_gain_by_region["roi"], [2.0])
    _assert_dist(stats.binding_gain_by_region["roi"], [1.0])
    _assert_dist(stats.func_cost_gain, [1.0, 2.0])
    _assert_dist(stats.binding_gain, [1.0, 1.0])
    _assert_empty_dist(stats.func_cost_gain_by_region["plateau"])
    _assert_empty_dist(stats.binding_gain_by_region["nonsyn"])


def test_next_run_chains():
    a1 = _run("s", solutions=[_sol("a", 10, 8, hits=[(1, 7)])], sweeps=["alpha"])
    a2 = _run(
        "s", alpha=2.0, solutions=[_sol("b", 9, 7, hits=[(1, 7)])], sweeps=["alpha"]
    )
    logp = _run("s", log_pos=True, solutions=[_sol("c", 1, 1)], sweeps=["alpha"])
    alpha_summaries = [_summary(r, ["alpha"]) for r in (a1, a2, logp)]
    attach_next_run_dominance([a1, a2, logp], alpha_summaries, w=500.0)
    assert alpha_summaries[0].dominance_alpha.n_by_region["hits"] == 1
    assert alpha_summaries[1].dominance_alpha.n_by_region["hits"] is None
    assert alpha_summaries[2].dominance_alpha.n_by_region["hits"] is None

    k50 = _run("s", k=50, log_pos=True, solutions=[_sol("d", 10, 8)], sweeps=["k"])
    k100 = _run("s", k=100, log_pos=True, solutions=[_sol("e", 9, 7)], sweeps=["k"])
    f_lo = _run(
        "t",
        fsm=1024,
        reduce=0.9375,
        log_pos=True,
        solutions=[_sol("f", 12, 9)],
        sweeps=["fsm_size"],
    )
    f_hi = _run(
        "t", fsm=2048, log_pos=True, solutions=[_sol("g", 8, 6)], sweeps=["fsm_size"]
    )
    k_sum = [_summary(k50, ["k"]), _summary(k100, ["k"])]
    f_sum = [_summary(f_lo, ["fsm_size"]), _summary(f_hi, ["fsm_size"])]
    attach_next_run_dominance([k50, k100, f_lo, f_hi], k_sum + f_sum, w=500.0)
    assert k_sum[0].dominance_k.n_by_region["roi"] == 1
    assert k_sum[1].dominance_k.n_by_region["roi"] is None
    assert f_sum[0].dominance_fsm.n_by_region["roi"] == 1
    assert f_sum[1].dominance_fsm.n_by_region["roi"] is None


def test_excel_columns_sort_and_headers(tmp_path: Path):
    from openpyxl import load_workbook

    from pareto_designer.views.experiment_report.xlsx_exporter import (
        export_experiment_xlsx,
    )

    summary_spec = design_run_table()
    alpha_spec = design_run_table("alpha")
    assert "n_dom_by_next" not in summary_spec.headers
    assert "func_cost_gain_mean" not in summary_spec.headers
    assert "binding_gain_mean" not in summary_spec.headers
    assert "n_dom_by_next" in alpha_spec.headers
    assert "func_cost_gain_mean" in alpha_spec.headers
    assert "binding_gain_mean" in alpha_spec.headers
    quality = next(g for g in alpha_spec.groups if g.title == "Solution quality")
    assert [sg.title for sg in quality.subgroups][-1] == "all"
    assert any(k.endswith(".n_global") for k in alpha_spec.keys)
    assert all(not k.startswith("dominance_k.") for k in alpha_spec.keys)
    sampling = summary_spec.groups[1]
    assert [g.title for g in sampling.subgroups] == ["low-cost preference", "k"]

    runs = [
        _summary(_run("s", alpha=2.0, solutions=[], sweeps=["alpha"]), ["alpha"]),
        _summary(
            _run(
                "s",
                alpha=0.5,
                log_pos=True,
                k=50,
                fsm=1024,
                solutions=[],
                sweeps=["alpha"],
            ),
            ["alpha"],
        ),
        _summary(_run("s", alpha=1.0, solutions=[], sweeps=["alpha"]), ["alpha"]),
    ]
    ordered = sort_design_runs(runs)
    assert [(r.log_pos, r.alpha) for r in ordered] == [
        (False, 1.0),
        (False, 2.0),
        (True, 0.5),
    ]

    run = _run("s", solutions=[_sol("a", 10, 8)], sweeps=["alpha"])
    out = tmp_path / "report.xlsx"
    export_experiment_xlsx(
        out,
        config=None,
        checklist=[],
        design_runs=build_design_run_summaries([run], config=None),
        solutions=[],
    )
    wb = load_workbook(out)
    assert wb["Summary"].freeze_panes == "A4"
    assert wb["Summary"]["A1"].value == "Sequence"
    sweep_leaf = next(
        wb["Sweep alpha"].iter_rows(min_row=3, max_row=3, values_only=True)
    )
    summary_leaf = next(wb["Summary"].iter_rows(min_row=3, max_row=3, values_only=True))
    assert "n_dom_by_next" in sweep_leaf
    assert "func_cost_gain_mean" in sweep_leaf
    assert "binding_gain_mean" in sweep_leaf
    assert "n_dom_by_next" not in summary_leaf
    assert "func_cost_gain_mean" not in summary_leaf
    assert "binding_gain_mean" not in summary_leaf
